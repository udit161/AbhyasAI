import pytest
import os
from pptx import Presentation
from pypdf import PdfWriter

from app.services.document_processor import DocumentProcessor, document_processor
from app.db.vector_store import vector_store
from app.models.vector_schema import VectorFilterQuery, ResourceType


def test_chunk_text_with_overlap():
    """
    Test text chunking with specified window size and overlap.
    """
    text = "Sample sentence for document parsing. " * 30  # ~1100 characters
    processor = DocumentProcessor(chunk_size=300, overlap=50)
    chunks = processor.chunk_text_with_overlap(text)

    assert len(chunks) >= 3
    assert len(chunks[0]) <= 300


def test_pdf_processing_and_ingestion(tmp_path):
    """
    Test creating a PDF, parsing pages, chunking with page_number metadata, and upserting into vector DB.
    """
    pdf_path = str(tmp_path / "sample_course.pdf")

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.add_blank_page(width=612, height=792)

    with open(pdf_path, "wb") as f:
        writer.write(f)

    doc_id = "test_pdf_101"
    title = "Machine Learning Architecture PDF"

    result = document_processor.process_and_ingest_pdf(
        file_path=pdf_path,
        doc_id=doc_id,
        title=title,
    )

    assert result["doc_id"] == doc_id
    assert result["doc_type"] == "pdf"


def test_ppt_processing_and_ingestion(tmp_path):
    """
    Test creating a PPTX presentation with python-pptx, extracting slide texts, tagging slide_index metadata, and upserting into vector DB.
    """
    ppt_path = str(tmp_path / "sample_presentation.pptx")

    prs = Presentation()
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "Slide 1: Gradient Descent Principles"

    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Slide 2: Vector Embedding and Similarity Search"

    prs.save(ppt_path)

    doc_id = "test_ppt_202"
    title = "AI Architecture Slides"

    result = document_processor.process_and_ingest_ppt(
        file_path=ppt_path,
        doc_id=doc_id,
        title=title,
    )

    assert result["doc_id"] == doc_id
    assert result["doc_type"] == "ppt"
    assert result["total_slides_extracted"] == 2

    # Query vector store specifically for slide_index = 2 and doc_id
    query_filter = VectorFilterQuery(
        query_text="Similarity Search",
        slide_index=2,
        allowed_resource_ids=[doc_id],
        resource_types=[ResourceType.PPT]
    )
    search_results = vector_store.query(query_filter)

    assert len(search_results) == 1

    assert search_results[0]["slide_index"] == 2
    assert "Similarity Search" in search_results[0]["content"]
