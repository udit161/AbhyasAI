import os
import argparse
import pdfplumber
from pptx import Presentation
from sentence_transformers import SentenceTransformer
import chromadb

# Text chunking constants
MAX_CHUNK_WORDS = 200
OVERLAP_WORDS = 50

def chunk_text(text: str, max_words: int, overlap: int) -> list[str]:
    """Simple word-based chunker with overlap."""
    words = text.split()
    chunks = []
    
    if not words:
        return chunks
        
    i = 0
    while i < len(words):
        chunk_words = words[i:i + max_words]
        chunks.append(" ".join(chunk_words))
        i += max_words - overlap
        
    return chunks

def process_pdf(file_path: str) -> list[dict]:
    """Extract text from PDF page by page."""
    print(f"Processing PDF: {file_path}")
    documents = []
    
    with pdfplumber.open(file_path) as pdf:
        for i, page in enumerate(pdf.pages):
            text = page.extract_text()
            if text:
                documents.append({
                    "text": text,
                    "page_number": i + 1
                })
    return documents

def process_pptx(file_path: str) -> list[dict]:
    """Extract text from PPTX slide by slide."""
    print(f"Processing PPTX: {file_path}")
    documents = []
    
    prs = Presentation(file_path)
    for i, slide in enumerate(prs.slides):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                slide_text.append(shape.text)
                
        if slide_text:
            documents.append({
                "text": " ".join(slide_text),
                "slide_number": i + 1
            })
    return documents

def ingest_document(file_path: str, doc_id: str, db_path: str = "./chroma_db"):
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == ".pdf":
        raw_docs = process_pdf(file_path)
        doc_type = "pdf"
        page_key = "page_number"
    elif ext in [".ppt", ".pptx"]:
        raw_docs = process_pptx(file_path)
        doc_type = "ppt"
        page_key = "slide_number"
    else:
        raise ValueError(f"Unsupported file type: {ext}")
        
    print(f"Extracted {len(raw_docs)} {page_key}s with text.")
    
    # Chunking
    final_chunks = []
    for doc in raw_docs:
        chunks = chunk_text(doc['text'], MAX_CHUNK_WORDS, OVERLAP_WORDS)
        for chunk in chunks:
            final_chunks.append({
                "text": chunk,
                page_key: doc[page_key]
            })
            
    print(f"Created {len(final_chunks)} overlapping chunks.")
    if not final_chunks:
        print("No text found. Exiting.")
        return
        
    # Generate Embeddings
    print("Loading embedding model (all-MiniLM-L6-v2)...")
    model = SentenceTransformer('all-MiniLM-L6-v2')
    
    texts = [c['text'] for c in final_chunks]
    print("Generating embeddings...")
    embeddings = model.encode(texts, show_progress_bar=True).tolist()
    
    # Upsert to ChromaDB
    print("Initializing ChromaDB...")
    client = chromadb.PersistentClient(path=db_path)
    collection = client.get_or_create_collection(name="course_materials")
    
    ids = [f"{doc_id}_chunk_{i}" for i in range(len(final_chunks))]
    metadatas = [{
        "doc_id": doc_id,
        "doc_type": doc_type,
        page_key: c[page_key] # dynamic: either 'page_number' or 'slide_number'
    } for c in final_chunks]
    
    print("Upserting to vector database...")
    collection.upsert(
        ids=ids,
        embeddings=embeddings,
        metadatas=metadatas,
        documents=texts
    )
    print(f"Successfully upserted {len(final_chunks)} chunks for {doc_id}!")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Ingest PDF/PPTX into vector db.")
    parser.add_argument("--input", required=True, help="Path to PDF or PPTX file")
    parser.add_argument("--doc_id", required=True, help="Unique identifier for the document")
    parser.add_argument("--db_path", default="./chroma_db", help="Path to store vector database")
    
    args = parser.parse_args()
    ingest_document(args.input, args.doc_id, args.db_path)
